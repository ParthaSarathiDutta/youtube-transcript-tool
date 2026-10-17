"""Word document generation for transcripts."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from transcript_service import TranscriptResult
from youtube_utils import PlaylistMetadata, VideoMetadata, sanitize_filename


@dataclass(frozen=True)
class DailyVideoDocument:
    video_id: str
    title: str
    url: str
    published_at_utc: datetime | None
    transcript: TranscriptResult


def write_single_video_document(
    metadata: VideoMetadata,
    transcript: TranscriptResult,
    output_dir: Path | str = ".",
) -> Path:
    """
    Create a .docx with:
    Title, URL, Transcript heading, body text.
    Timestamps are stored on TranscriptResult.segments but not shown by default.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = f"{sanitize_filename(metadata.title)}_transcript.docx"
    output_path = output_dir / filename

    doc = Document()
    doc.add_heading(metadata.title, level=0)
    doc.add_paragraph(metadata.url)
    doc.add_heading("Transcript", level=1)
    doc.add_paragraph(transcript.display_text)

    doc.save(output_path)
    return output_path


def write_playlist_document(
    playlist: PlaylistMetadata,
    items: list[tuple[VideoMetadata, TranscriptResult]],
    output_dir: Path | str = ".",
) -> Path:
    """Create one .docx containing every video transcript in playlist order."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = f"{sanitize_filename(playlist.title)}_transcripts.docx"
    output_path = output_dir / filename

    doc = Document()
    doc.add_heading(playlist.title, level=0)

    for index, (video, transcript) in enumerate(items, start=1):
        if index > 1:
            doc.add_paragraph()
            doc.add_paragraph("---")
            paragraph = doc.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)

        doc.add_heading(f"Video {index}: {video.title}", level=2)
        doc.add_paragraph(f"YouTube URL: {video.url}")
        doc.add_heading("Transcript", level=3)
        doc.add_paragraph(transcript.display_text)

    doc.save(output_path)
    return output_path


def write_daily_combined_document(
    digest_title: str,
    items: list[DailyVideoDocument],
    output_path: Path | str,
) -> Path:
    """Combined digest .docx with full transcript text (no summary substitution)."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(digest_title, level=0)

    success_items = [item for item in items if item.transcript.available]
    for index, item in enumerate(success_items, start=1):
        if index > 1:
            paragraph = doc.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)

        doc.add_heading(f"Video {index}: {item.title}", level=1)
        doc.add_paragraph(f"YouTube URL: {item.url}")
        if item.published_at_utc:
            doc.add_paragraph(f"Published (UTC): {item.published_at_utc.isoformat()}")
        doc.add_heading("Transcript", level=2)
        doc.add_paragraph(item.transcript.display_text)

    doc.save(output_path)
    return output_path
